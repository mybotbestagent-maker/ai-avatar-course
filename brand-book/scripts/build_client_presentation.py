#!/usr/bin/env python3
"""Build Gold Hands Painting client presentation PPTX."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "Gold-Hands-Painting-Client-Presentation.pptx"
LOGO = ROOT / "logo" / "chosen" / "logo-transparent-master.png"

ORANGE = RGBColor(0xF1, 0x5A, 0x24)
NAVY = RGBColor(0x1A, 0x27, 0x44)
NAVY_DEEP = RGBColor(0x0E, 0x16, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF3, 0xF1, 0xEC)
MUTED = RGBColor(0x5B, 0x65, 0x78)


def set_run(run, size=18, bold=False, color=NAVY, font="Montserrat"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color):
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    fill.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_bar(slide, top=True, color=ORANGE):
    y = 0 if top else Inches(7.35)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, Inches(13.333), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=MUTED):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.level = 0
        for run in p.runs:
            set_run(run, size=size, color=color)
        if not p.runs:
            run = p.add_run()
            run.text = "• " + item
            set_run(run, size=size, color=color)
        p.space_after = Pt(8)
    return box


def card(slide, left, top, w, h, title, body, dark=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x25, 0x34, 0x56) if dark else WHITE
    shape.line.color.rgb = RGBColor(0x3A, 0x4A, 0x6A) if dark else RGBColor(0xE5, 0xE1, 0xD8)
    add_text(slide, left + Inches(0.2), top + Inches(0.18), w - Inches(0.4), Inches(0.4), title, size=15, bold=True, color=ORANGE if dark else NAVY)
    add_text(slide, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.7), body, size=13, color=WHITE if dark else MUTED)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY_DEEP)
    add_bar(s, True, ORANGE)
    add_bar(s, False, ORANGE)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(0.7), Inches(0.7), height=Inches(1.8))
    add_text(s, Inches(0.7), Inches(2.7), Inches(11), Inches(0.4), "CLIENT PRESENTATION · MIAMI", 14, True, ORANGE)
    add_text(s, Inches(0.7), Inches(3.2), Inches(11), Inches(1.8), "Your home,\npainted the right way.", 44, True, WHITE)
    add_text(s, Inches(0.7), Inches(5.3), Inches(10), Inches(0.8), "Interior & exterior · Free written estimates · 5-year warranty path", 18, False, WHITE)
    add_text(s, Inches(0.7), Inches(6.3), Inches(10), Inches(0.4), "(786) 788-8714  ·  goldhandsmia.com", 16, True, ORANGE)

    # 2 Who
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "WHO WE ARE", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.8), "Local Miami painters you can trust", 32, True, NAVY)
    add_text(s, Inches(0.7), Inches(1.7), Inches(11), Inches(0.7), "Gold Hands Painting — painting division of Gold Hands Miami. Insured local crews across Miami-Dade & Broward.", 16, False, MUTED)
    cards = [
        ("Free estimate", "Written & itemized — no pressure."),
        ("Insured crews", "Background-checked painters."),
        ("5-year warranty", "Transferable path on premium 2-coat."),
        ("EN / ES / RU", "Project manager support."),
    ]
    for i, (t, b) in enumerate(cards):
        card(s, Inches(0.7 + i * 3.1), Inches(2.8), Inches(2.9), Inches(2.2), t, b)

    # 3 Services
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "WHAT WE DO", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Full painting services", 32, True, NAVY)
    card(s, Inches(0.7), Inches(1.9), Inches(5.8), Inches(2.4), "Interior", "Walls, ceilings, trim, doors & baseboards. Accent walls & full-home repaints. Condo/HOA scheduling & COI. Daily cleanup. Typical: 2–3 days.")
    card(s, Inches(6.8), Inches(1.9), Inches(5.8), Inches(2.4), "Exterior & stucco", "Stucco, siding, trim, fascia. Elastomeric / hurricane-ready options. Pressure washing. Colors for Florida light. Typical: 3–5 days.")
    card(s, Inches(0.7), Inches(4.5), Inches(5.8), Inches(1.8), "Color consultation", "Finishes that survive Miami sun, salt air, and HOA rules — before paint day.")
    card(s, Inches(6.8), Inches(4.5), Inches(5.8), Inches(1.8), "Also available", "Popcorn ceiling removal, cabinet painting consults, punch-list touch-ups.")

    # 4 Process
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY_DEEP)
    add_bar(s, True, ORANGE)
    add_bar(s, False, ORANGE)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "HOW WE WORK", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Clear process. No surprises.", 32, True, WHITE)
    steps = [
        ("1", "Call / request", "Tell us the scope. We reply fast."),
        ("2", "On-site measure", "Walk the job & note condition."),
        ("3", "Written quote", "Itemized price before start."),
        ("4", "Prep & paint", "Protect, repair, prime, coat."),
        ("5", "Walkthrough", "Final check + warranty docs."),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = Inches(0.5 + i * 2.5)
        card(s, x, Inches(2.2), Inches(2.35), Inches(3.5), f"{n}. {t}", b, dark=True)

    # 5 Day by day
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "ON YOUR JOB", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "What happens from day one", 32, True, NAVY)
    card(s, Inches(0.7), Inches(1.9), Inches(5.8), Inches(2.3), "Before we start", "Confirm scope, colors, schedule in writing. Typical start 48–72h. HOA/condo COI when required.")
    card(s, Inches(6.8), Inches(1.9), Inches(5.8), Inches(2.3), "On paint days", "Protect furniture & floors. Patch, cut clean edges. Primer + premium coats. Daily cleanup.")
    card(s, Inches(0.7), Inches(4.4), Inches(5.8), Inches(1.9), "Materials", "Quality systems for Florida humidity — value, premium, or hurricane elastomeric.")
    card(s, Inches(6.8), Inches(4.4), Inches(5.8), Inches(1.9), "Your PM", "One contact updates you. Support in English, Spanish, Russian.")

    # 6 Prep
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "PREP", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Why our finish lasts", 32, True, NAVY)
    add_text(s, Inches(0.7), Inches(1.7), Inches(11), Inches(0.5), "The finish only looks good if prep is right.", 18, False, MUTED)
    card(s, Inches(0.7), Inches(2.5), Inches(3.8), Inches(3.2), "Protect", "Masking, drop cloths, plastic — floors, furniture, fixtures covered.")
    card(s, Inches(4.75), Inches(2.5), Inches(3.8), Inches(3.2), "Repair", "Patch holes, fix cracks, sand, caulk. Exterior wash & crack repair when needed.")
    card(s, Inches(8.8), Inches(2.5), Inches(3.8), Inches(3.2), "Prime & coat", "Proper primer where needed, then clean premium coats with sharp lines.")

    # 7 Pricing
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY_DEEP)
    add_bar(s, True, ORANGE)
    add_bar(s, False, ORANGE)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "PRICING", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Clear price before we start", 32, True, WHITE)
    card(s, Inches(0.7), Inches(2.0), Inches(5.8), Inches(2.0), "Free written estimate", "On-site measure + itemized quote. Usually within 24–48 hours.", True)
    card(s, Inches(6.8), Inches(2.0), Inches(5.8), Inches(2.0), "No surprise fees", "Agreed scope = agreed price. Changes only with your OK.", True)
    card(s, Inches(0.7), Inches(4.3), Inches(5.8), Inches(2.0), "AI budget preview", "Rough range online, then free on-site quote confirms exact price.", True)
    card(s, Inches(6.8), Inches(4.3), Inches(5.8), Inches(2.0), "Special", "$0.99/sq ft when painted area > 3,000 sq ft (Standard materials).", True)

    # 8 Warranty
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "PROTECTION", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Warranty & insurance", 32, True, NAVY)
    card(s, Inches(0.7), Inches(2.0), Inches(5.8), Inches(4.0), "5-year transferable warranty path", "Peeling, blistering, flaking & excessive fading on premium 2-coat with standard prep. Can transfer to a buyer. Not storm/substrate failure.")
    card(s, Inches(6.8), Inches(2.0), Inches(5.8), Inches(4.0), "Insured & bonded", "General Liability. Background-checked crews. COI for HOA/condo. We make workmanship right.")

    # 9 Area + promises
    s = prs.slides.add_slide(blank)
    add_bg(s, PAPER)
    add_bar(s, True, ORANGE)
    add_bar(s, False, NAVY)
    add_text(s, Inches(0.7), Inches(0.45), Inches(11), Inches(0.35), "WHY GOLD HANDS", 13, True, ORANGE)
    add_text(s, Inches(0.7), Inches(0.9), Inches(11), Inches(0.7), "Promises we keep", 32, True, NAVY)
    promises = [
        ("Upfront pricing", "Price for agreed scope before we begin."),
        ("Clean jobsite", "Protect your home and clean as we go."),
        ("On-time updates", "Your PM keeps you informed."),
        ("Make-right", "Not happy? We come back and fix it."),
        ("Local", "Miami-Dade & Broward — 130+ ZIP codes."),
        ("Fast start", "Often within 48–72 hours."),
    ]
    for i, (t, b) in enumerate(promises):
        r, c = divmod(i, 3)
        card(s, Inches(0.7 + c * 4.1), Inches(1.9 + r * 2.4), Inches(3.9), Inches(2.15), t, b)

    # 10 Close
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY_DEEP)
    add_bar(s, True, ORANGE)
    add_bar(s, False, ORANGE)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(0.7), Inches(0.6), height=Inches(1.6))
    add_text(s, Inches(0.7), Inches(2.5), Inches(11), Inches(0.4), "NEXT STEP", 14, True, ORANGE)
    add_text(s, Inches(0.7), Inches(3.0), Inches(11), Inches(1.2), "Let’s paint your home.", 44, True, WHITE)
    add_text(s, Inches(0.7), Inches(4.4), Inches(11), Inches(0.5), "Free on-site written estimate. Clear scope. Clean finish.", 18, False, WHITE)
    add_text(s, Inches(0.7), Inches(5.3), Inches(11), Inches(0.6), "(786) 788-8714", 36, True, ORANGE)
    add_text(s, Inches(0.7), Inches(6.2), Inches(11), Inches(0.4), "goldhandsmia.com  ·  @_gold_hands_  ·  Miami-Dade & Broward", 16, True, WHITE)

    prs.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    build()
